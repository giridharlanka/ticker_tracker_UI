"""Tests for Analysis and Portfolio Signals workbook sheets."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from ticker_tracker.analysis.base import (
    FundamentalsResult,
    LLMAnalysis,
    SignalLabel,
    TechnicalResult,
)
from ticker_tracker.report_builder import _fill_rgb, build_portfolio_workbook


def _minimal_metadata() -> dict:
    return {
        "run_timestamp_utc": "2026-01-01T00:00:00+00:00",
        "fx_source": "frankfurter",
        "fx_rates": [],
        "finance_source_by_ticker": {"AAPL": "yahoo"},
        "price_fetch_failed": [],
        "fx_unavailable_tickers": [],
    }


def _holding(ticker: str = "AAPL") -> dict:
    return {
        "ticker": ticker,
        "shares": 10.0,
        "report_ccy": "USD",
        "cost_per_share_purchase": 100.0,
        "cost_basis_purchase": 1000.0,
        "price_per_share_purchase": 120.0,
        "current_value_purchase": 1200.0,
        "gain_loss_purchase": 200.0,
        "gain_loss_pct_purchase": 20.0,
        "cost_basis_base": 1000.0,
        "current_value_base": 1200.0,
    }


def _fundamentals(ticker: str = "AAPL", *, available: bool = True) -> FundamentalsResult:
    return FundamentalsResult(
        ticker=ticker,
        source="yahoo",
        fetched_at=datetime.now(UTC),
        fundamentals_available=available,
        pe_ratio=28.0,
        pb_ratio=5.0,
        net_margin=0.22,
        revenue_growth_yoy=0.08,
        debt_to_equity=50.0,
        analyst_recommendation="buy",
        target_price=200.0,
        last_earnings_surprise_pct=3.5,
    )


def _technicals(ticker: str = "AAPL") -> TechnicalResult:
    return TechnicalResult(
        ticker=ticker,
        analysis_date=datetime.now(UTC),
        period_days=365,
        rsi_14=55.0,
        macd_histogram=0.12,
        price_vs_ema200_pct=4.5,
        signals=[
            SignalLabel("Above 200-day EMA", bullish=True),
            SignalLabel("MACD Positive", bullish=True),
        ],
    )


def _analysis(ticker: str, signal: str) -> LLMAnalysis:
    return LLMAnalysis(
        ticker=ticker,
        model="qwen2.5:7b",
        signal=signal,
        confidence="High",
        strengths=["s1", "s2", "s3"],
        risks=["r1", "r2", "r3"],
        summary="Strong outlook.",
        generated_at=datetime.now(UTC),
        prompt_tokens_approx=500,
        llm_available=True,
    )


def _summary() -> dict:
    return {
        "total_cost_basis_base": 1000.0,
        "total_current_value_base": 1200.0,
        "total_gain_loss_base": 200.0,
        "total_return_pct": 20.0,
        "holding_count": 1,
        "distinct_ticker_count": 1,
        "totals_purchase_cost_by_ccy": {"USD": 1000.0},
        "totals_purchase_value_by_ccy": {"USD": 1200.0},
        "totals_purchase_gl_by_ccy": {"USD": 200.0},
    }


def test_workbook_with_analysis_has_five_sheets(tmp_path: Path) -> None:
    out = tmp_path / "report.xlsx"
    build_portfolio_workbook(
        out,
        base_currency="USD",
        holdings_rows=[_holding()],
        summary=_summary(),
        metadata=_minimal_metadata(),
        fundamentals={"AAPL": _fundamentals()},
        technicals={"AAPL": _technicals()},
        analyses={"AAPL": _analysis("AAPL", "BUY")},
        portfolio_summary={"commentary": "Balanced.", "top_action": "Hold winners."},
        llm_available=True,
    )
    wb = load_workbook(out)
    assert wb.sheetnames == [
        "Holdings",
        "Summary",
        "Metadata",
        "Analysis",
        "Portfolio Signals",
    ]


def test_buy_signal_cell_green_fill(tmp_path: Path) -> None:
    out = tmp_path / "buy.xlsx"
    build_portfolio_workbook(
        out,
        base_currency="USD",
        holdings_rows=[_holding()],
        summary=_summary(),
        metadata=_minimal_metadata(),
        fundamentals={"AAPL": _fundamentals()},
        technicals={"AAPL": _technicals()},
        analyses={"AAPL": _analysis("AAPL", "BUY")},
        portfolio_summary={"commentary": "x", "top_action": "y"},
        llm_available=True,
    )
    ws = load_workbook(out)["Analysis"]
    signal_cell = ws.cell(2, 2)
    assert signal_cell.value == "BUY"
    assert _fill_rgb(signal_cell.fill) == "C6EFCE"


def test_sell_signal_cell_red_fill(tmp_path: Path) -> None:
    out = tmp_path / "sell.xlsx"
    build_portfolio_workbook(
        out,
        base_currency="USD",
        holdings_rows=[_holding()],
        summary=_summary(),
        metadata=_minimal_metadata(),
        fundamentals={"AAPL": _fundamentals()},
        technicals={"AAPL": _technicals()},
        analyses={"AAPL": _analysis("AAPL", "SELL")},
        portfolio_summary={"commentary": "x", "top_action": "y"},
        llm_available=True,
    )
    ws = load_workbook(out)["Analysis"]
    signal_cell = ws.cell(2, 2)
    assert signal_cell.value == "SELL"
    assert _fill_rgb(signal_cell.fill) == "FFC7CE"


def test_without_analysis_args_only_three_sheets(tmp_path: Path) -> None:
    out = tmp_path / "legacy.xlsx"
    build_portfolio_workbook(
        out,
        base_currency="USD",
        holdings_rows=[_holding()],
        summary=_summary(),
        metadata=_minimal_metadata(),
    )
    wb = load_workbook(out)
    assert wb.sheetnames == ["Holdings", "Summary", "Metadata"]


def test_sparse_data_quality_in_analysis_sheet(tmp_path: Path) -> None:
    out = tmp_path / "sparse.xlsx"
    build_portfolio_workbook(
        out,
        base_currency="USD",
        holdings_rows=[_holding("D05.SI")],
        summary=_summary(),
        metadata=_minimal_metadata(),
        fundamentals={"D05.SI": _fundamentals("D05.SI", available=False)},
        technicals={"D05.SI": _technicals("D05.SI")},
        analyses=None,
        llm_available=False,
    )
    ws = load_workbook(out)["Analysis"]
    assert "SPARSE" in str(ws.cell(2, 18).value)
