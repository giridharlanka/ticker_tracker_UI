"""Portfolio XLSX report (Holdings, Summary, Metadata, optional Analysis)."""

from __future__ import annotations

from collections import Counter
from collections.abc import Callable, Mapping
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ticker_tracker.analysis.base import (
    FundamentalsResult,
    LLMAnalysis,
    SignalLabel,
    TechnicalResult,
)
from ticker_tracker.calculator import purchase_amount_lines_by_ccy
from ticker_tracker.fx.base import FXRate

_HEADER_FONT = Font(bold=True)
_POS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_NEG_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_PCT_POS_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_PCT_NEG_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

_ANALYSIS_HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
_ANALYSIS_HEADER_FONT = Font(bold=True, color="FFFFFF")
_SIGNAL_FILLS = {
    "BUY": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "HOLD": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "SELL": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "INSUFFICIENT DATA": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
    "LLM OFFLINE": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
}
_RSI_ALERT_FILL = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
_FONT_GREEN = Font(color="006100")
_FONT_RED = Font(color="9C0006")
_FONT_GREY = Font(color="808080")
_HEATMAP_BULL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_HEATMAP_BEAR = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

_ANALYSIS_HEADERS = [
    "Ticker",
    "Signal",
    "Confidence",
    "RSI(14)",
    "MACD Histogram",
    "vs EMA200 (%)",
    "P/E",
    "P/B",
    "Net Margin (%)",
    "Rev Growth YoY (%)",
    "Debt/Equity",
    "Analyst Rec",
    "Target Price",
    "Earnings Surprise (%)",
    "Strengths",
    "Risks",
    "LLM Summary",
    "Data Quality",
]

_HEATMAP_COLUMNS: list[tuple[str, Callable[[SignalLabel], bool]]] = [
    ("RSI", lambda s: "RSI" in s.name),
    ("MACD", lambda s: "MACD" in s.name),
    ("vs EMA200", lambda s: "EMA" in s.name and "200" in s.name),
    ("Golden/Death Cross", lambda s: "Cross" in s.name and "EMA" in s.name),
    ("BB", lambda s: "Bollinger" in s.name),
    ("Volume", lambda s: "Volume" in s.name),
]


def _sheet_filename_timestamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def default_workbook_filename() -> str:
    return f"Ticker_summary_{_sheet_filename_timestamp()}.xlsx"


def _pct_fill(value: Any) -> PatternFill | None:
    if not isinstance(value, int | float):
        return None
    if value > 0:
        return _PCT_POS_FILL
    if value < 0:
        return _PCT_NEG_FILL
    return None


def _fill_rgb(fill: PatternFill | None) -> str | None:
    if fill is None or fill.fill_type != "solid":
        return None
    rgb = fill.start_color.rgb
    return str(rgb)[-6:].upper() if rgb else None


def _autofit_worksheet(ws: Worksheet, *, min_width: int = 12, max_width: int = 60) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells:
            val = cell.value
            if val is None:
                continue
            max_len = max(max_len, len(str(val)))
        ws.column_dimensions[letter].width = max(min_width, min(max_width, max_len + 2))


def _style_analysis_header(ws: Worksheet, row: int, ncol: int) -> None:
    for col in range(1, ncol + 1):
        c = ws.cell(row, col)
        c.font = _ANALYSIS_HEADER_FONT
        c.fill = _ANALYSIS_HEADER_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row + 1, 1)


def _heatmap_signals(technicals: TechnicalResult | None) -> dict[str, list[SignalLabel]]:
    grouped: dict[str, list[SignalLabel]] = {name: [] for name, _ in _HEATMAP_COLUMNS}
    if technicals is None:
        return grouped
    for sig in technicals.signals:
        for col_name, pred in _HEATMAP_COLUMNS:
            if pred(sig):
                grouped[col_name].append(sig)
    return grouped


def _add_analysis_sheet(
    wb: Workbook,
    *,
    holdings_rows: list[dict[str, Any]],
    fundamentals: dict[str, FundamentalsResult],
    technicals: dict[str, TechnicalResult],
    analyses: dict[str, LLMAnalysis] | None,
    llm_available: bool,
) -> None:
    ws = wb.create_sheet("Analysis")
    ncol = len(_ANALYSIS_HEADERS)
    for col, title in enumerate(_ANALYSIS_HEADERS, start=1):
        ws.cell(1, col, title)
    _style_analysis_header(ws, 1, ncol)

    col_idx = {h: i + 1 for i, h in enumerate(_ANALYSIS_HEADERS)}

    for row_i, holding in enumerate(holdings_rows, start=2):
        ticker = str(holding.get("ticker") or "")
        fund = fundamentals.get(ticker)
        tech = technicals.get(ticker)
        analysis = (analyses or {}).get(ticker)

        ws.cell(row_i, col_idx["Ticker"], ticker)

        if analysis is not None and analysis.llm_available:
            signal = analysis.signal
            ws.cell(row_i, col_idx["Signal"], signal)
            ws.cell(row_i, col_idx["Confidence"], analysis.confidence)
            ws.cell(row_i, col_idx["Strengths"], " | ".join(analysis.strengths))
            ws.cell(row_i, col_idx["Risks"], " | ".join(analysis.risks))
            ws.cell(row_i, col_idx["LLM Summary"], analysis.summary)
        else:
            signal = "LLM OFFLINE"
            ws.cell(row_i, col_idx["Signal"], signal)
            ws.cell(row_i, col_idx["LLM Summary"], "LLM offline — per-ticker analysis unavailable.")

        sig_fill = _SIGNAL_FILLS.get(signal)
        if sig_fill:
            ws.cell(row_i, col_idx["Signal"]).fill = sig_fill

        if tech is not None:
            if tech.rsi_14 is not None:
                rsi_cell = ws.cell(row_i, col_idx["RSI(14)"], tech.rsi_14)
                if tech.rsi_14 < 30 or tech.rsi_14 > 70:
                    rsi_cell.fill = _RSI_ALERT_FILL
                    rsi_cell.font = Font(bold=True)
            if tech.macd_histogram is not None:
                ws.cell(row_i, col_idx["MACD Histogram"], tech.macd_histogram)
            if tech.price_vs_ema200_pct is not None:
                ema_cell = ws.cell(row_i, col_idx["vs EMA200 (%)"], tech.price_vs_ema200_pct / 100)
                ema_cell.number_format = "0.0%"
                if tech.price_vs_ema200_pct > 0:
                    ema_cell.font = _FONT_GREEN
                elif tech.price_vs_ema200_pct < 0:
                    ema_cell.font = _FONT_RED

        if fund is not None:
            if fund.pe_ratio is not None:
                ws.cell(row_i, col_idx["P/E"], fund.pe_ratio)
            if fund.pb_ratio is not None:
                ws.cell(row_i, col_idx["P/B"], fund.pb_ratio)
            if fund.net_margin is not None:
                c = ws.cell(row_i, col_idx["Net Margin (%)"], fund.net_margin)
                c.number_format = "0.0%"
                if fund.net_margin < 0:
                    c.font = _FONT_RED
            if fund.revenue_growth_yoy is not None:
                c = ws.cell(row_i, col_idx["Rev Growth YoY (%)"], fund.revenue_growth_yoy)
                c.number_format = "0.0%"
                if fund.revenue_growth_yoy < 0:
                    c.font = _FONT_RED
            if fund.debt_to_equity is not None:
                ws.cell(row_i, col_idx["Debt/Equity"], fund.debt_to_equity)
            if fund.analyst_recommendation is not None:
                ws.cell(row_i, col_idx["Analyst Rec"], fund.analyst_recommendation)
            if fund.target_price is not None:
                ws.cell(row_i, col_idx["Target Price"], fund.target_price)
            if fund.last_earnings_surprise_pct is not None:
                c = ws.cell(
                    row_i,
                    col_idx["Earnings Surprise (%)"],
                    fund.last_earnings_surprise_pct / 100,
                )
                c.number_format = "0.0%"
                if fund.last_earnings_surprise_pct < 0:
                    c.font = _FONT_RED
            quality = "FULL" if fund.fundamentals_available else f"SPARSE — {ticker}"
            q_cell = ws.cell(row_i, col_idx["Data Quality"], quality)
            if not fund.fundamentals_available:
                q_cell.font = _FONT_GREY

    _autofit_worksheet(ws)


def _add_portfolio_signals_sheet(
    wb: Workbook,
    *,
    holdings_rows: list[dict[str, Any]],
    technicals: dict[str, TechnicalResult],
    analyses: dict[str, LLMAnalysis] | None,
    portfolio_summary: Mapping[str, str] | None,
    llm_offline_hint: str,
    llm_available: bool,
) -> None:
    ws = wb.create_sheet("Portfolio Signals")
    analyses_map = analyses or {}
    signals = [a.signal for a in analyses_map.values() if a.llm_available]
    n = len(holdings_rows)
    counts = Counter(signals)
    total_llm = len(signals) or 1

    def pct(sig: str) -> str:
        if not signals:
            return "0%"
        return f"{counts.get(sig, 0) / total_llm * 100:.0f}%"

    conf_counter = Counter(a.confidence for a in analyses_map.values() if a.llm_available)
    conf_breakdown = ", ".join(f"{k}: {v}" for k, v in sorted(conf_counter.items())) or "—"

    summary_rows = [
        ("Total Holdings", n),
        ("BUY signals", f"{counts.get('BUY', 0)} ({pct('BUY')})"),
        ("HOLD signals", f"{counts.get('HOLD', 0)} ({pct('HOLD')})"),
        ("SELL signals", f"{counts.get('SELL', 0)} ({pct('SELL')})"),
        ("Avg Confidence", conf_breakdown),
        ("LLM Available", "Yes" if llm_available else "No (Offline)"),
    ]
    for i, (label, value) in enumerate(summary_rows, start=1):
        ws.cell(i, 1, label).font = _HEADER_FONT
        ws.cell(i, 2, value)

    ws.cell(8, 1, "Portfolio Commentary (AI-generated)").font = Font(bold=True, size=12)
    ws.merge_cells("A8:B8")
    if llm_available and portfolio_summary:
        ws.cell(9, 1, "Commentary")
        ws.cell(9, 2, portfolio_summary.get("commentary", ""))
        ws.cell(10, 1, "Top Action")
        ws.cell(10, 2, portfolio_summary.get("top_action", ""))
    else:
        ws.cell(
            9,
            1,
            f"LLM offline — commentary unavailable. {llm_offline_hint}",
        )
        ws.merge_cells("A9:B9")

    heatmap_row = 14
    heat_headers = ["Ticker", *[name for name, _ in _HEATMAP_COLUMNS]]
    for col, title in enumerate(heat_headers, start=1):
        c = ws.cell(heatmap_row, col, title)
        c.font = _ANALYSIS_HEADER_FONT
        c.fill = _ANALYSIS_HEADER_FILL
    ws.freeze_panes = ws.cell(heatmap_row + 1, 1)

    for offset, holding in enumerate(holdings_rows, start=1):
        row = heatmap_row + offset
        ticker = str(holding.get("ticker") or "")
        ws.cell(row, 1, ticker)
        tech = technicals.get(ticker)
        grouped = _heatmap_signals(tech)
        for col_offset, (col_name, _) in enumerate(_HEATMAP_COLUMNS, start=2):
            matches = grouped[col_name]
            if not matches:
                continue
            label = matches[0].name
            cell = ws.cell(row, col_offset, label)
            if matches[0].bullish:
                cell.fill = _HEATMAP_BULL
            else:
                cell.fill = _HEATMAP_BEAR

    _autofit_worksheet(ws)


def _mixed_purchase_currencies(holdings_rows: list[dict[str, Any]]) -> bool:
    """True if rows use more than one *report_ccy* (purchase-currency totals would mix units)."""
    seen: set[str] = set()
    for h in holdings_rows:
        raw = h.get("report_ccy")
        if raw is None:
            continue
        s = str(raw).strip().upper()
        if not s or s == "—":
            continue
        seen.add(s)
    return len(seen) > 1


def build_portfolio_workbook(
    path: str | Path,
    *,
    base_currency: str,
    holdings_rows: list[dict[str, Any]],
    summary: Mapping[str, Any],
    metadata: Mapping[str, Any],
    fundamentals: dict[str, FundamentalsResult] | None = None,
    technicals: dict[str, TechnicalResult] | None = None,
    analyses: dict[str, LLMAnalysis] | None = None,
    portfolio_summary: Mapping[str, str] | None = None,
    llm_offline_hint: str = "Ensure the analysis LLM is configured and reachable.",
    llm_available: bool | None = None,
) -> Path:
    """
    Write workbook to *path* (Holdings, Summary, Metadata; optional Analysis sheets).

    *holdings_rows* must supply keys aligned with the engine (see ``engine.py``).
    When *fundamentals*, *technicals*, and *analyses* are all ``None``, analysis sheets
    are omitted (legacy three-sheet workbook).
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_h = wb.active
    assert ws_h is not None
    ws_h.title = "Holdings"
    b = base_currency.upper()

    h1 = [
        "Ticker",
        "Shares",
        "Cost/sh (purch.)",
        "CCY",
        "Cost basis (purch.)",
        "Price/sh (purch.)",
        "Latest value (purch.)",
        "G/L (purch.)",
        "G/L %",
    ]
    for col, title in enumerate(h1, start=1):
        c = ws_h.cell(row=1, column=col, value=title)
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    pct_col = 9
    gl_amt_col = 8
    last_data = len(holdings_rows) + 1
    for r, holding in enumerate(holdings_rows, start=2):
        ws_h.cell(r, 1, holding.get("ticker"))
        ws_h.cell(r, 2, holding.get("shares"))
        ws_h.cell(r, 3, holding.get("cost_per_share_purchase"))
        ws_h.cell(r, 4, holding.get("report_ccy"))
        ws_h.cell(r, 5, holding.get("cost_basis_purchase"))
        ws_h.cell(r, 6, holding.get("price_per_share_purchase"))
        ws_h.cell(r, 7, holding.get("current_value_purchase"))
        gl_p = holding.get("gain_loss_purchase")
        ws_h.cell(r, gl_amt_col, gl_p)
        pct_val = holding.get("gain_loss_pct_purchase")
        ws_h.cell(r, pct_col, pct_val)
        if isinstance(gl_p, int | float):
            if gl_p > 0:
                ws_h.cell(r, gl_amt_col).fill = _POS_FILL
            elif gl_p < 0:
                ws_h.cell(r, gl_amt_col).fill = _NEG_FILL
        pf = _pct_fill(pct_val)
        if pf:
            ws_h.cell(r, pct_col).fill = pf

    if not _mixed_purchase_currencies(holdings_rows):
        total_row = last_data + 1
        ws_h.cell(total_row, 1, "TOTAL").font = _HEADER_FONT
        if last_data >= 2:
            ws_h.cell(total_row, 2, f"=SUM(B2:B{last_data})").font = _HEADER_FONT
            ws_h.cell(total_row, 5, f"=SUM(E2:E{last_data})").font = _HEADER_FONT
            ws_h.cell(total_row, 7, f"=SUM(G2:G{last_data})").font = _HEADER_FONT
            ws_h.cell(total_row, 8, f"=SUM(H2:H{last_data})").font = _HEADER_FONT

    for col in range(1, len(h1) + 1):
        ws_h.column_dimensions[get_column_letter(col)].width = 16

    ws_s = wb.create_sheet("Summary")
    ws_s["A1"] = f"Base currency: {b}"
    ws_s["A1"].font = Font(bold=True, size=14)
    ws_s["A3"] = "Metric"
    ws_s["B3"] = "Base"
    ws_s["C3"] = "Purchased"
    for c in (1, 2, 3):
        ws_s.cell(3, c).font = _HEADER_FONT

    r = 4
    bu = b.upper()

    def fmt_base_cell(v: Any) -> str | float:
        if isinstance(v, bool) or not isinstance(v, int | float):
            return v if v is not None else "—"
        return f"{float(v):,.2f} {bu}"

    cost_map: dict[str, float] = dict(summary.get("totals_purchase_cost_by_ccy") or {})
    val_map: dict[str, float] = dict(summary.get("totals_purchase_value_by_ccy") or {})
    gl_map: dict[str, float] = dict(summary.get("totals_purchase_gl_by_ccy") or {})

    if not cost_map and summary.get("totals_purchase_cost_formatted"):
        ws_s.cell(r, 1, "Total invested")
        ws_s.cell(r, 2, fmt_base_cell(summary.get("total_cost_basis_base")))
        ws_s.cell(r, 3, summary.get("totals_purchase_cost_formatted") or "—")
        r += 1
        ws_s.cell(r, 1, "Current value")
        ws_s.cell(r, 2, fmt_base_cell(summary.get("total_current_value_base")))
        ws_s.cell(r, 3, summary.get("totals_purchase_value_formatted") or "—")
        r += 1
        ws_s.cell(r, 1, "Gain / loss")
        ws_s.cell(r, 2, fmt_base_cell(summary.get("total_gain_loss_base")))
        ws_s.cell(r, 3, summary.get("totals_purchase_gl_formatted") or "—")
        r += 1
    else:
        for label, base_key, pmap in (
            ("Total invested", "total_cost_basis_base", cost_map),
            ("Current value", "total_current_value_base", val_map),
            ("Gain / loss", "total_gain_loss_base", gl_map),
        ):
            lines = purchase_amount_lines_by_ccy(pmap)
            ws_s.cell(r, 1, label)
            ws_s.cell(r, 2, fmt_base_cell(summary.get(base_key)))
            ws_s.cell(r, 3, lines[0])
            r += 1
            for ln in lines[1:]:
                ws_s.cell(r, 1, "")
                ws_s.cell(r, 2, "")
                ws_s.cell(r, 3, ln)
                r += 1

    tr = summary.get("total_return_pct")
    ws_s.cell(r, 1, "Total return %")
    ws_s.cell(r, 2, tr)
    ws_s.cell(r, 3, tr)
    r += 1

    ws_s.cell(r, 1, "Holdings count (distinct tickers)")
    ws_s.cell(r, 2, summary.get("distinct_ticker_count", summary.get("holding_count")))
    ws_s.cell(r, 3, "—")
    r += 1

    row = r + 1
    for note in summary.get("assumption_notes") or []:
        ws_s.cell(row, 1, note)
        ws_s.cell(row, 1).alignment = Alignment(wrap_text=True)
        row += 1

    ws_m = wb.create_sheet("Metadata")
    ws_m["A1"] = "Run timestamp (UTC)"
    ws_m["B1"] = metadata.get("run_timestamp_utc", "")
    ws_m["A2"] = "Base currency"
    ws_m["B2"] = b
    ws_m["A3"] = "FX source"
    ws_m["B3"] = metadata.get("fx_source", "")
    r = 5
    ws_m.cell(r, 1, "FX rates used").font = _HEADER_FONT
    r += 1
    for c, title in enumerate(["From", "To", "Rate", "Fetched At (UTC)", "Source"], start=1):
        ws_m.cell(r, c, title).font = _HEADER_FONT
    r += 1
    for fx in metadata.get("fx_rates") or []:
        if isinstance(fx, FXRate):
            ws_m.cell(r, 1, fx.from_currency)
            ws_m.cell(r, 2, fx.to_currency)
            ws_m.cell(r, 3, fx.rate)
            ws_m.cell(r, 4, fx.fetched_at.astimezone(UTC).isoformat())
            ws_m.cell(r, 5, fx.source)
        elif isinstance(fx, dict):
            ws_m.cell(r, 1, fx.get("from_currency"))
            ws_m.cell(r, 2, fx.get("to_currency"))
            ws_m.cell(r, 3, fx.get("rate"))
            ws_m.cell(r, 4, fx.get("fetched_at"))
            ws_m.cell(r, 5, fx.get("source"))
        r += 1

    r += 1
    ws_m.cell(r, 1, "Finance source per ticker").font = _HEADER_FONT
    r += 1
    ws_m.cell(r, 1, "Ticker").font = _HEADER_FONT
    ws_m.cell(r, 2, "Source").font = _HEADER_FONT
    r += 1
    for t, src in sorted((metadata.get("finance_source_by_ticker") or {}).items()):
        ws_m.cell(r, 1, t)
        ws_m.cell(r, 2, src)
        r += 1

    r += 1
    ws_m.cell(r, 1, "Price fetch failed (tickers)").font = _HEADER_FONT
    r += 1
    ws_m.cell(r, 1, ", ".join(metadata.get("price_fetch_failed") or []) or "—")
    r += 2
    ws_m.cell(r, 1, "FX rate unavailable (tickers)").font = _HEADER_FONT
    r += 1
    ws_m.cell(r, 1, ", ".join(metadata.get("fx_unavailable_tickers") or []) or "—")

    include_analysis = fundamentals is not None or technicals is not None or analyses is not None
    if include_analysis and fundamentals is not None and technicals is not None:
        ollama_ok = llm_available if llm_available is not None else bool(analyses)
        _add_analysis_sheet(
            wb,
            holdings_rows=holdings_rows,
            fundamentals=fundamentals,
            technicals=technicals,
            analyses=analyses,
            llm_available=ollama_ok,
        )
        _add_portfolio_signals_sheet(
            wb,
            holdings_rows=holdings_rows,
            technicals=technicals,
            analyses=analyses,
            portfolio_summary=portfolio_summary,
            llm_offline_hint=llm_offline_hint,
            llm_available=ollama_ok,
        )

    wb.save(path)
    return path


def build_portfolio_html_analysis_section(
    *,
    holdings_rows: list[dict[str, Any]],
    fundamentals: dict[str, FundamentalsResult],
    technicals: dict[str, TechnicalResult],
    analyses: dict[str, LLMAnalysis] | None,
    portfolio_summary: Mapping[str, str] | None,
    llm_offline_hint: str,
    llm_available: bool,
) -> str:
    """HTML fragment for Analysis and Portfolio Signals (appended to HTML report)."""
    import html as html_mod

    parts: list[str] = ['<h3>Analysis</h3><table border="1" cellpadding="4" cellspacing="0">']
    header_cells = "".join(f"<th>{html_mod.escape(h)}</th>" for h in _ANALYSIS_HEADERS)
    parts.append(f"<tr>{header_cells}</tr>")
    for holding in holdings_rows:
        ticker = str(holding.get("ticker") or "")
        fund = fundamentals.get(ticker)
        tech = technicals.get(ticker)
        analysis = (analyses or {}).get(ticker)
        signal = analysis.signal if analysis and analysis.llm_available else "LLM OFFLINE"
        cells = [
            ticker,
            signal,
            analysis.confidence if analysis and analysis.llm_available else "—",
            f"{tech.rsi_14:.1f}" if tech and tech.rsi_14 is not None else "—",
            f"{tech.macd_histogram:.4f}" if tech and tech.macd_histogram is not None else "—",
            (
                f"{tech.price_vs_ema200_pct:+.1f}%"
                if tech and tech.price_vs_ema200_pct is not None
                else "—"
            ),
            f"{fund.pe_ratio:.2f}" if fund and fund.pe_ratio is not None else "—",
            f"{fund.pb_ratio:.2f}" if fund and fund.pb_ratio is not None else "—",
            (f"{fund.net_margin * 100:.1f}%" if fund and fund.net_margin is not None else "—"),
            (
                f"{fund.revenue_growth_yoy * 100:.1f}%"
                if fund and fund.revenue_growth_yoy is not None
                else "—"
            ),
            (f"{fund.debt_to_equity:.2f}" if fund and fund.debt_to_equity is not None else "—"),
            fund.analyst_recommendation if fund and fund.analyst_recommendation else "—",
            f"{fund.target_price:.2f}" if fund and fund.target_price is not None else "—",
            (
                f"{fund.last_earnings_surprise_pct:+.1f}%"
                if fund and fund.last_earnings_surprise_pct is not None
                else "—"
            ),
            (" | ".join(analysis.strengths) if analysis and analysis.llm_available else "—"),
            (" | ".join(analysis.risks) if analysis and analysis.llm_available else "—"),
            (
                analysis.summary
                if analysis and analysis.llm_available
                else "LLM offline — per-ticker analysis unavailable."
            ),
            ("FULL" if fund and fund.fundamentals_available else f"SPARSE — {ticker}"),
        ]
        parts.append(
            "<tr>" + "".join(f"<td>{html_mod.escape(str(c))}</td>" for c in cells) + "</tr>"
        )
    parts.append("</table>")

    parts.append("<h3>Portfolio Signals</h3>")
    if llm_available and portfolio_summary:
        parts.append(
            f"<p><b>Commentary:</b> {html_mod.escape(portfolio_summary.get('commentary', ''))}</p>"
        )
        parts.append(
            f"<p><b>Top action:</b> {html_mod.escape(portfolio_summary.get('top_action', ''))}</p>"
        )
    else:
        parts.append(
            "<p><i>LLM offline — commentary unavailable. "
            f"{html_mod.escape(llm_offline_hint)}</i></p>"
        )

    return "".join(parts)
